"""Front matter: cover, title, declaration, certificate, abstract,
acknowledgement, contents, lists of figures/tables/acronyms."""

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from rlib import (CONTENT_W, FONT, chapter, figure, frontmatter_heading, page_break,
                  para, rich, section, subsection, table)

# ---------------------------------------------------------------------------
# Everything the student must personalise is collected here. Nothing in this
# file invents a registration number, a guide or an HOD: those are facts about
# a real institution and a real person, and a report is an official document.
# ---------------------------------------------------------------------------
STUDENT = "NISHANT E S"
REGNO = "<<REGISTER NUMBER>>"
GUIDE = "<<GUIDE NAME>>"
GUIDE_DESIG = "<<DESIGNATION>>"
HOD = "<<HOD NAME>>"
PROGRAM = "COMPUTER SCIENCE AND ENGINEERING"
PROGRAM_TC = "Computer Science and Engineering"
MONTH_YEAR = "November, 2025"
TITLE = ("THERAPYTRACE: A LONGITUDINAL LINGUISTIC BIOMARKER "
         "MONITORING SYSTEM FOR PSYCHOTHERAPY")
TITLE_TC = ("TherapyTrace: A Longitudinal Linguistic Biomarker "
            "Monitoring System for Psychotherapy")
LOGO = "fig/vit_logo.jpg"


def _title_block(doc, cover=True):
    para(doc, "", after=24)
    para(doc, "A project report on", size=12, italic=True, align="center", after=18)
    para(doc, TITLE, size=20, bold=True, align="center", after=20, line=1.15)
    para(doc, "Submitted in partial fulfillment for the award of the degree of",
         size=14, italic=True, align="center", after=16)
    para(doc, f"Bachelor of Technology in {PROGRAM_TC}",
         size=22, bold=True, align="center", after=20, line=1.15)
    para(doc, "by", size=14, italic=True, align="center", after=14)
    para(doc, f"{STUDENT} ({REGNO})", size=16, bold=True, align="center", after=28)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    p.add_run().add_picture(LOGO, width=Cm(8.8))

    para(doc, "SCHOOL OF COMPUTER SCIENCE AND ENGINEERING",
         size=16, bold=True, align="center", after=10)
    para(doc, MONTH_YEAR, size=12, align="center", after=0)


def cover_and_title(doc):
    _title_block(doc, cover=True)
    page_break(doc)
    _title_block(doc, cover=False)


def declaration(doc):
    # No page break here: the section break that switches numbering to roman
    # has already started a fresh page, and adding another leaves a blank one.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    p.add_run().add_picture(LOGO, width=Cm(8.8))

    frontmatter_heading(doc, "DECLARATION")
    rich(doc, [
        ("I hereby declare that the thesis entitled ", False, False),
        (f"“{TITLE_TC}”", True, False),
        (f" submitted by {STUDENT} ({REGNO}), for the award of the degree of "
         f"Bachelor of Technology in {PROGRAM_TC}, Vellore Institute of "
         f"Technology, Chennai is a record of bonafide work carried out by me "
         f"under the supervision of {GUIDE}.", False, False),
    ], after=12)
    para(doc, "I further declare that the work reported in this thesis has not "
              "been submitted and will not be submitted, either in part or in "
              "full, for the award of any other degree or diploma in this "
              "institute or any other institute or university.", after=36)
    para(doc, "Place: Chennai", align="left", after=8)
    para(doc, "Date:", align="left", after=30)
    para(doc, "Signature of the Candidate", bold=True, align="right", after=0)


def certificate(doc):
    page_break(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(LOGO, width=Cm(8.8))
    para(doc, "School of Computer Science and Engineering",
         size=14, bold=True, align="center", after=16)

    frontmatter_heading(doc, "CERTIFICATE", underline=False)
    rich(doc, [
        ("This is to certify that the report entitled ", False, False),
        (f"“{TITLE_TC}”", True, False),
        (f" is prepared and submitted by {STUDENT} ({REGNO}) to Vellore "
         f"Institute of Technology, Chennai, in partial fulfillment of the "
         f"requirement for the award of the degree of ", False, False),
        (f"Bachelor of Technology in {PROGRAM_TC}", True, False),
        (" is a bonafide record carried out under my guidance. The project "
         "fulfills the requirements as per the regulations of this University "
         "and in my opinion meets the necessary standards for submission. The "
         "contents of this report have not been submitted and will not be "
         "submitted either in part or in full, for the award of any other "
         "degree or diploma and the same is certified.", False, False),
    ], after=26)

    para(doc, "Signature of the Guide:", align="left", after=8)
    para(doc, f"Name: {GUIDE}", align="left", after=8)
    para(doc, "Date:", align="left", after=26)

    t = doc.add_table(rows=3, cols=2)
    t.autofit = False
    data = [("Signature of the Examiner", "Signature of the Examiner"),
            ("Name:", "Name:"), ("Date:", "Date:")]
    for ri, (a, b) in enumerate(data):
        for ci, val in enumerate((a, b)):
            cell = t.rows[ri].cells[ci]
            cell.text = ""
            cell.width = Cm(CONTENT_W.cm / 2)
            pp = cell.paragraphs[0]
            pp.paragraph_format.line_spacing = 1.5
            r = pp.add_run(val)
            r.font.name, r.font.size = FONT, Pt(12)
    para(doc, "", after=20)

    para(doc, "Approved by the Head of Department,", align="left", after=6)
    para(doc, f"B.Tech. {PROGRAM}", bold=True, align="left", after=14)
    para(doc, f"Name: {HOD}", align="left", after=8)
    para(doc, "Date:", align="left", after=18)
    para(doc, "(Seal of the School)", bold=True, align="left", after=0)


def abstract(doc):
    page_break(doc)
    frontmatter_heading(doc, "ABSTRACT")
    para(doc,
         "Psychotherapy outcome is assessed almost entirely through "
         "self-report questionnaires administered weeks apart. Between those "
         "measurements the clinician has only memory and session notes, and "
         "the language of the session itself — the richest signal available — "
         "is discarded. This project builds TherapyTrace, a full-stack system "
         "that measures how a client’s language changes across successive "
         "therapy sessions and presents that change as a trajectory rather "
         "than as a single score.")
    para(doc,
         "The system extracts five therapeutic process dimensions from client "
         "speech — self-agency, future orientation, emotional granularity, "
         "problem ownership and reflection depth — and scores each session "
         "against the client’s own earlier sessions rather than against a "
         "population norm. This within-person or idiographic framing is the "
         "central methodological commitment of the work: a person who speaks "
         "in a characteristically passive way is not compared with people who "
         "do not, but with how that same person spoke last month. Small-sample "
         "variance is stabilised by shrinking each client’s standard deviation "
         "toward a cohort prior, and the resulting Therapeutic Progress Index "
         "is tracked over time with trend estimation and change-point "
         "detection.")
    para(doc,
         "A supervised layer trained on AnnoMI, a corpus of 133 "
         "expert-annotated motivational interviewing demonstrations, provides "
         "an independent second signal. Seven classifiers were compared under "
         "a protocol in which splits are grouped by conversation rather than "
         "by utterance, since an utterance-level split allows a model to "
         "recognise a conversation it has memorised. The comparison produced a "
         "deliberately reported negative result: on the client talk-type task "
         "a paired t-test cannot separate the best learner from any of the "
         "other six, and the spread across all seven is smaller than the "
         "fold-to-fold variation of a single model.")
    para(doc,
         "That negative result motivated the principal contribution. Rather "
         "than pursuing a better classifier, the work measures how classifier "
         "error propagates into the quantity a clinician would actually read. "
         "A classifier at 0.49 macro-F1 was shown to produce a session-level "
         "change-talk ratio biased upward by 0.25, a bias that is systematic "
         "and therefore does not diminish as sessions grow longer, and which "
         "compresses a true range of 0.10–0.90 into a displayed range of "
         "0.63–0.85. A prevalence correction restores fidelity from a slope "
         "of 0.27 to 0.80. The textbook correction was found to clip at "
         "realistic session lengths, so a Bayesian posterior over the true "
         "ratio is used instead, which remains within range by construction "
         "and yields a credible interval the interface displays.")
    para(doc,
         "The system is implemented as a FastAPI backend with a React "
         "frontend, comprising twenty-five modules and sixty-nine automated "
         "tests. It is positioned as decision support for reflection and "
         "supervision, not as a diagnostic instrument, and its limitations — "
         "chiefly the absence of a real longitudinal clinical corpus — are "
         "stated explicitly rather than minimised.", after=0)


def acknowledgement(doc):
    page_break(doc)
    frontmatter_heading(doc, "ACKNOWLEDGEMENT", underline=False)
    rich(doc, [
        ("It is my pleasure to express with deep sense of gratitude to ", False, False),
        (GUIDE, True, False),
        (f", {GUIDE_DESIG}, School of Computer Science and Engineering, "
         "Vellore Institute of Technology, Chennai, for the constant guidance, "
         "continual encouragement and understanding; more than all, for "
         "teaching me patience in my endeavor. My association is not confined "
         "to academics only, but it has been a great opportunity on my part to "
         "work with an intellectual and expert in the field of clinical "
         "natural language processing.", False, False),
    ])
    para(doc,
         "It is with gratitude that I would like to extend my thanks to the "
         "visionary leader Dr. G. Viswanathan our Honorable Chancellor, Mr. "
         "Sankar Viswanathan, Dr. Sekar Viswanathan, Dr. G V Selvam Vice "
         "Presidents, Dr. Sandhya Pentareddy, Executive Director, Ms. "
         "Kadhambari S. Viswanathan, Assistant Vice-President, Dr. V. S. "
         "Kanchana Bhaaskaran Vice-Chancellor, Dr. T. Thyagarajan Pro-Vice "
         "Chancellor, VIT Chennai and Dr. P. K. Manoharan, Additional "
         "Registrar for providing an exceptional working environment and "
         "inspiring all of us during the tenure of the course.")
    para(doc,
         "Special mention to Dr. Viswanathan V, Dean, Dr. Nithyanandam P, "
         "Associate Dean, Dr. Suganya G, Associate Dean, Dr. Sweetlin "
         "Hemalatha C, Associate Dean, School of Computer Science and "
         "Engineering, Vellore Institute of Technology, Chennai for spending "
         "their valuable time and efforts in sharing their knowledge and for "
         "helping us in every aspect.")
    rich(doc, [
        ("In jubilant state, I express ingeniously my whole-hearted thanks to ",
         False, False),
        (HOD, True, False),
        (", Head of the Department, B.Tech. ", False, False),
        (PROGRAM, True, False),
        (" and the Project Coordinators for their valuable support and "
         "encouragement to take up and complete the thesis.", False, False),
    ])
    para(doc,
         "My sincere thanks to all the faculties and staff at Vellore "
         "Institute of Technology, Chennai who helped me acquire the requisite "
         "knowledge. I would like to thank my parents for their support. It is "
         "indeed a pleasure to thank my friends who encouraged me to take up "
         "and complete this task.", after=30)
    para(doc, "Place: Chennai", align="left", after=8)
    para(doc, "Date:", align="left", after=24)
    para(doc, STUDENT, bold=True, align="right", after=0)
